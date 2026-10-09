# -*- coding: utf-8 -*-
"""把 OCR 得到的逐页 Markdown 转成 Word，并**附加页码信息**。

页码处理两种模式：
- 默认：每页前插入一行 **【第 N 页】** 标记（可导航、可搜索、可复制），文本连续排版；
- `--page-breaks`：每个 PDF 页另起一个 Word 页，便于与纸本逐页对照。

排版复用 `tools/md2docx_unified.py`（中文论文版式：宋体小四、1.5 倍行距、首行缩进、页脚 PAGE 域）。

用法：
    python -B _ocr_md_to_docx.py <输入.OCR.md> [输出.docx] [--page-breaks]
"""
import argparse
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)
sys.path.insert(0, str(Path(__file__).resolve().parent / "tools"))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

import md2docx_unified as M  # noqa: E402  （复用版式与 Markdown 渲染）

PAGE_SPLIT = re.compile(r"^--- 第 (\d+) 页 ---$", re.MULTILINE)


def add_page_label(doc, number: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"【第 {number} 页】")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("md")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--page-breaks", action="store_true", help="每个 PDF 页另起一个 Word 页")
    a = ap.parse_args()

    md = Path(a.md).resolve()
    out = Path(a.out).resolve() if a.out else md.with_suffix(".docx")
    raw = md.read_text(encoding="utf-8", errors="replace")

    # 去掉脚本自动生成的说明头（保留其中的标题行作为文档标题）
    title = md.stem
    body = raw
    if raw.startswith("# "):
        first_line, _, rest = raw.partition("\n")
        title = first_line[2:].strip() or title
        body = rest
    parts = PAGE_SPLIT.split(body)
    # parts: [前言, 页号1, 内容1, 页号2, 内容2, ...]
    pairs: list[tuple[str, str]] = []
    for i in range(1, len(parts) - 1, 2):
        pairs.append((parts[i], parts[i + 1]))
    if not pairs:  # 没有页标记：整体当一页处理
        pairs = [("1", body)]

    doc = Document()
    M.setup_styles(doc)
    M.setup_page(doc, page_number=True)

    h = doc.add_heading("", level=1)
    M.add_runs(h, title)
    M.fmt_paragraph(h, "heading")

    total_chars = 0
    for idx, (number, page_text) in enumerate(pairs):
        if a.page_breaks and idx > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        add_page_label(doc, number)
        blocks = M.parse_markdown(page_text, None)
        M.render_blocks(doc, blocks)
        total_chars += sum(len(b.get("text", "")) for b in blocks)

    doc.save(str(out))
    print(f"完成：{out}")
    print(f"  页数 {len(pairs)}｜正文字符 {total_chars:,}｜大小 {out.stat().st_size / 1024 / 1024:.1f} MB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
